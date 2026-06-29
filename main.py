from __future__ import annotations

import csv
import io
import os
import re
import shutil
import tempfile
import zipfile
from collections import OrderedDict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import fitz  # PyMuPDF
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles

CM_TO_PT = 72 / 2.54
SIZE_MAP = {
    "10x8": (10.0, 8.0),
    "10x10": (10.0, 10.0),
    "10x15": (10.0, 15.0),
}

# Broad SellerSKU pattern. Real labels may include letters, numbers, and hyphens.
SKU_TOKEN_PATTERN = re.compile(r"[A-Z0-9][A-Z0-9_-]{3,}")
FBA_SHIPMENT_PATTERN = re.compile(r"FBA[A-Z0-9]{8,}")
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]+')

# PDF text operators supported by this version.
PDF_LITERAL_TJ_RE = re.compile(rb"\((?:\\.|[^\\)])*\)\s*Tj")
PDF_ARRAY_TJ_RE = re.compile(rb"\[(?:\\.|[^\]])*\]\s*TJ", re.S)
PDF_TEXT_TOKEN_RE = re.compile(rb"\((?:\\.|[^\\)])*\)|<\s*[0-9A-Fa-f\s]+\s*>", re.S)


def safe_name(name: str) -> str:
    name = INVALID_FILENAME_CHARS.sub("-", str(name or "").strip())
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name or "UNKNOWN-SKU"


def cm_to_pt(v: float) -> float:
    return v * CM_TO_PT


def pdf_literal_unescape(data: bytes) -> bytes:
    """Small PDF literal string unescaper, enough for label text matching."""
    out = bytearray()
    i = 0
    while i < len(data):
        b = data[i]
        if b == 0x5C and i + 1 < len(data):  # backslash
            n = data[i + 1]
            mapping = {
                ord("n"): b"\n",
                ord("r"): b"\r",
                ord("t"): b"\t",
                ord("b"): b"\b",
                ord("f"): b"\f",
                ord("("): b"(",
                ord(")"): b")",
                ord("\\"): b"\\",
            }
            if n in mapping:
                out.extend(mapping[n])
                i += 2
                continue
            if 48 <= n <= 55:  # octal escape \ddd
                j = i + 1
                oct_digits = []
                while j < len(data) and len(oct_digits) < 3 and 48 <= data[j] <= 55:
                    oct_digits.append(chr(data[j]))
                    j += 1
                try:
                    out.append(int("".join(oct_digits), 8) & 0xFF)
                    i = j
                    continue
                except Exception:
                    pass
            # Line continuation or unknown escape: keep escaped byte.
            out.append(n)
            i += 2
            continue
        out.append(b)
        i += 1
    return bytes(out)


def pdf_literal_escape_ascii(text: str) -> bytes:
    raw = text.encode("ascii")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def pdf_hex_to_bytes(token: bytes) -> bytes:
    inner = token.strip()[1:-1]
    inner = re.sub(rb"\s+", b"", inner)
    if len(inner) % 2 == 1:
        inner += b"0"
    try:
        return bytes.fromhex(inner.decode("ascii"))
    except Exception:
        return b""


def decode_pdf_text_token(token: bytes) -> str:
    token = token.strip()
    if token.startswith(b"("):
        raw = pdf_literal_unescape(token[1:-1])
    elif token.startswith(b"<"):
        raw = pdf_hex_to_bytes(token)
    else:
        return ""

    if raw.startswith(b"\xfe\xff"):
        try:
            return raw[2:].decode("utf-16-be", errors="ignore")
        except Exception:
            return ""

    if b"\x00" in raw:
        try:
            return raw.decode("utf-16-be", errors="ignore")
        except Exception:
            pass

    try:
        return raw.decode("latin1", errors="ignore")
    except Exception:
        return ""


def build_fba_replacement_from_array(array_token: bytes) -> bytes:
    """
    Replace a TJ array containing FBA:xxxx with native FBA only.
    For per-character arrays, preserve the original F/B/A glyph tokens and kerning.
    For UTF-16BE hex arrays, use a UTF-16BE hex FBA token.
    """
    content_start = array_token.find(b"[")
    content_end = array_token.rfind(b"]")
    if content_start < 0 or content_end <= content_start:
        return b"[(FBA)] TJ"

    content = array_token[content_start + 1 : content_end]
    string_matches = list(PDF_TEXT_TOKEN_RE.finditer(content))
    decoded = "".join(decode_pdf_text_token(m.group(0)) for m in string_matches)

    if not (decoded.startswith("FBA:") or decoded.startswith("FBA：")):
        return array_token

    # If F, B, A are separate tokens, keep the original glyph tokens up to A.
    char_count = 0
    cut_pos = None
    for m in string_matches:
        txt = decode_pdf_text_token(m.group(0))
        char_count += len(txt)
        if char_count == 3:
            cut_pos = m.end()
            break
        if char_count > 3:
            break

    if cut_pos is not None:
        return b"[" + content[:cut_pos] + b"] TJ"

    first = string_matches[0].group(0) if string_matches else b""
    if first.strip().startswith(b"<"):
        # UTF-16BE hex for FBA.
        return b"[<004600420041>] TJ"

    return b"[(FBA)] TJ"


def replace_fba_line_text_in_streams(doc: fitz.Document) -> int:
    """
    True PDF content-stream replacement.

    Replaces destination line text only when it starts with:
      FBA: ...
      FBA：...

    Supported source encodings:
      (FBA: company) Tj
      [(F)-0.000(B)-0.000(A)-0.000(:)...] TJ
      [<004600420041003a...>] TJ

    This function never draws white boxes, so it will not cover the ship-from
    name/address, destination title, warehouse code, or warehouse address.
    """
    total = 0

    for page in doc:
        for xref in page.get_contents() or []:
            try:
                stream = doc.xref_stream(xref)
            except Exception:
                continue

            changed = False

            def repl_literal_tj(m: re.Match[bytes]) -> bytes:
                nonlocal changed, total
                token = m.group(0)
                start = token.find(b"(")
                end = token.rfind(b")")
                if start < 0 or end <= start:
                    return token

                inner = token[start + 1 : end]
                plain = pdf_literal_unescape(inner)
                plain_str = plain.decode("latin1", errors="ignore")

                if plain_str.startswith("FBA:") or plain_str.startswith("FBA："):
                    changed = True
                    total += 1
                    suffix = token[end + 1 :]  # includes optional whitespace and Tj
                    return b"(" + pdf_literal_escape_ascii("FBA") + b")" + suffix

                if plain.startswith(b"\x00F\x00B\x00A\x00:") or plain.startswith(b"\x00F\x00B\x00A\xff\x1a"):
                    changed = True
                    total += 1
                    suffix = token[end + 1 :]
                    return b"<004600420041>" + suffix

                return token

            def repl_array_tj(m: re.Match[bytes]) -> bytes:
                nonlocal changed, total
                token = m.group(0)
                content_start = token.find(b"[")
                content_end = token.rfind(b"]")
                if content_start < 0 or content_end <= content_start:
                    return token

                content = token[content_start + 1 : content_end]
                string_matches = list(PDF_TEXT_TOKEN_RE.finditer(content))
                decoded = "".join(decode_pdf_text_token(mm.group(0)) for mm in string_matches)

                if decoded.startswith("FBA:") or decoded.startswith("FBA："):
                    changed = True
                    total += 1
                    return build_fba_replacement_from_array(token)

                return token

            new_stream = PDF_LITERAL_TJ_RE.sub(repl_literal_tj, stream)
            new_stream = PDF_ARRAY_TJ_RE.sub(repl_array_tj, new_stream)

            if changed and new_stream != stream:
                doc.update_stream(xref, new_stream)

    return total


def extract_lines(page: fitz.Page) -> List[str]:
    text = page.get_text("text") or ""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if lines:
        return lines

    d = page.get_text("dict")
    rows: List[Tuple[float, float, str]] = []
    for b in d.get("blocks", []):
        for line in b.get("lines", []):
            y = min((s.get("bbox", [0, 0, 0, 0])[1] for s in line.get("spans", [])), default=0)
            x = min((s.get("bbox", [0, 0, 0, 0])[0] for s in line.get("spans", [])), default=0)
            t = "".join(s.get("text", "") for s in line.get("spans", [])).strip()
            if t:
                rows.append((y, x, t))
    return [t for _, _, t in sorted(rows)]


def normalize_sku_candidate(text: str, allow_short: bool = False) -> str:
    t = re.sub(r"\s+", "", text.strip().upper())
    t = t.strip("：:;,.，。")
    if not t:
        return ""
    if t.startswith("FBA") or FBA_SHIPMENT_PATTERN.match(t):
        return ""
    if any(word in t for word in {"GUANGDONG", "SHENZHEN", "PERRYSBURG", "FREMONT", "中国", "美国", "深圳"}):
        return ""
    if t in {"SINGLESKU", "SKU", "MADEINCHINA"}:
        return ""
    if "数量" in t or t.startswith("QTY") or t.startswith("COUNT"):
        return ""

    m = SKU_TOKEN_PATTERN.search(t)
    if not m:
        return ""
    candidate = m.group(0)
    if "-" not in candidate:
        if allow_short:
            if len(candidate) < 4 or not re.search(r"[A-Z]", candidate) or not re.search(r"\d", candidate):
                return ""
        elif len(candidate) < 8 or not re.search(r"[A-Z]", candidate) or not re.search(r"\d", candidate):
            return ""
    if len(candidate) < 5:
        return ""
    return candidate


def extract_sku_and_qty(page: fitz.Page) -> Tuple[str, int]:
    lines = extract_lines(page)
    sku = ""
    qty = 1

    for i, line in enumerate(lines):
        if "Single SKU" in line:
            after = line.split("Single SKU", 1)[1].strip()
            direct = normalize_sku_candidate(after, allow_short=True)
            if direct:
                sku = direct
                break
            for candidate in lines[i + 1 : i + 10]:
                if "数量" in candidate or candidate.lower().startswith("made in"):
                    break
                cand = normalize_sku_candidate(candidate, allow_short=True)
                if cand:
                    sku = cand
                    break
            break

    if not sku:
        for line in lines:
            cand = normalize_sku_candidate(line)
            if cand:
                sku = cand
                break

    for i, line in enumerate(lines):
        if "数量" in line:
            m = re.search(r"数量\s*[:：]?\s*(\d+)", line)
            if m:
                qty = int(m.group(1))
            elif i + 1 < len(lines):
                m2 = re.search(r"\b(\d+)\b", lines[i + 1])
                if m2:
                    qty = int(m2.group(1))
            break

    return (sku or "UNKNOWN-SKU", qty)


def source_has_bottom_warning(page: fitz.Page) -> bool:
    try:
        return "请不要遮住此标签" in (page.get_text("text") or "")
    except Exception:
        return False


def make_output_page(
    src_doc: fitz.Document,
    page_index: int,
    size_key: str,
    add_made: bool,
    made_font_size: float,
) -> fitz.Document:
    width_cm, height_cm = SIZE_MAP[size_key]
    target_w = cm_to_pt(width_cm)
    target_h = cm_to_pt(height_cm)

    tmp = fitz.open()
    tmp.insert_pdf(src_doc, from_page=page_index, to_page=page_index)
    replace_fba_line_text_in_streams(tmp)

    page = tmp[0]
    original_page = src_doc[page_index]
    src_w, src_h = page.rect.width, page.rect.height
    scale = target_w / src_w
    clip_h = min(src_h, target_h / scale)
    clip = fitz.Rect(0, 0, src_w, clip_h)

    out = fitz.open()
    out_page = out.new_page(width=target_w, height=target_h)
    out_page.show_pdf_page(
        fitz.Rect(0, 0, target_w, target_h),
        tmp,
        0,
        clip=clip,
        keep_proportion=False,
    )

    if add_made:
        made_text = "Made In China"
        made_font_size = float(made_font_size or 8.0)
        made_offset_cm = 0.25 if source_has_bottom_warning(original_page) else 0.5
        made_y = min(target_h - 2.5, target_h - cm_to_pt(made_offset_cm))
        text_w = fitz.get_text_length(made_text, fontname="helv", fontsize=made_font_size)
        made_x = max(0, (target_w - text_w) / 2)
        out_page.insert_text(
            (made_x, made_y),
            made_text,
            fontsize=made_font_size,
            fontname="helv",
            color=(0, 0, 0),
            overlay=True,
        )

    tmp.close()
    return out


def collect_pdfs(input_path: Path, work_dir: Path) -> List[Path]:
    if input_path.suffix.lower() == ".pdf":
        return [input_path]
    if input_path.suffix.lower() == ".zip":
        extract_dir = work_dir / "input_zip"
        extract_dir.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(input_path, "r") as zf:
            zf.extractall(extract_dir)
        return sorted(extract_dir.rglob("*.pdf"))
    raise ValueError("只支持 PDF 或 ZIP 文件")


def load_sku_mapping(mapping_path: Optional[Path]) -> Dict[str, str]:
    if not mapping_path or not mapping_path.exists() or mapping_path.stat().st_size == 0:
        return {}

    def norm_key(v: object) -> str:
        return re.sub(r"\s+", "", str(v or "").strip().upper())

    def norm_header(v: object) -> str:
        return re.sub(r"[\s_\-]+", "", str(v or "").strip().lower())

    rows: List[List[object]] = []
    suffix = mapping_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        from openpyxl import load_workbook

        wb = load_workbook(str(mapping_path), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            values = list(row or [])
            if any(str(c or "").strip() for c in values):
                rows.append(values)
        wb.close()
    elif suffix == ".csv":
        raw = mapping_path.read_bytes()
        text = raw.decode("utf-8-sig", errors="ignore")
        for row in csv.reader(io.StringIO(text)):
            if any(str(c or "").strip() for c in row):
                rows.append(row)
    else:
        return {}

    if not rows:
        return {}

    header = [norm_header(c) for c in rows[0]]
    seller_names = {"sellersku", "sellerskuid", "sellsku", "sellerskucode", "sellerskuno", "msku", "sku", "卖家sku", "销售sku"}
    warehouse_names = {"映射仓库sku", "仓库sku", "仓库skuid", "warehousesku", "warehouse", "whsku", "仓库编码", "内部sku", "本地sku"}
    seller_idx = next((i for i, h in enumerate(header) if h in seller_names or ("seller" in h and "sku" in h)), None)
    warehouse_idx = next((i for i, h in enumerate(header) if h in warehouse_names or ("warehouse" in h and "sku" in h) or ("仓库" in h and "sku" in h)), None)

    if seller_idx is None or warehouse_idx is None:
        seller_idx, warehouse_idx = 0, 1
        data_rows = rows
        if len(rows[0]) >= 2 and ("sku" in header[0] or "仓库" in header[1]):
            data_rows = rows[1:]
    else:
        data_rows = rows[1:]

    mapping: Dict[str, str] = {}
    for row in data_rows:
        if len(row) <= max(seller_idx, warehouse_idx):
            continue
        seller = norm_key(row[seller_idx])
        warehouse = str(row[warehouse_idx] or "").strip()
        if seller and warehouse:
            mapping[seller] = warehouse
    return mapping


def mapped_sku_for_naming(seller_sku: str, naming_mode: str, sku_mapping: Dict[str, str]) -> str:
    if naming_mode == "warehouse":
        key = re.sub(r"\s+", "", str(seller_sku or "").strip().upper())
        return sku_mapping.get(key) or seller_sku or "UNKNOWN-SKU"
    return seller_sku or "UNKNOWN-SKU"


def write_group_pdf(doc: fitz.Document, out_dir: Path, base: str, used_bases: Dict[str, int]) -> Path:
    safe_base = safe_name(base)
    count = used_bases.get(safe_base, 0) + 1
    used_bases[safe_base] = count
    final_base = safe_base if count == 1 else f"{safe_base}-{count}"
    folder = out_dir / final_base
    folder.mkdir(parents=True, exist_ok=True)
    out_pdf = folder / f"{final_base}.pdf"
    doc.save(str(out_pdf), deflate=True, garbage=4)
    return out_pdf


def process_file(
    input_path: Path,
    output_zip: Path,
    size_key: str,
    add_made: bool,
    made_font_size: float,
    group_by_sku: bool,
    naming_mode: str = "seller",
    mapping_path: Optional[Path] = None,
) -> Dict:
    if size_key not in SIZE_MAP:
        raise ValueError("尺寸参数无效")
    if naming_mode not in {"seller", "warehouse"}:
        naming_mode = "seller"

    sku_mapping = load_sku_mapping(mapping_path) if naming_mode == "warehouse" else {}

    with tempfile.TemporaryDirectory() as td:
        work_dir = Path(td)
        pdf_paths = collect_pdfs(input_path, work_dir)
        if not pdf_paths:
            raise ValueError("没有找到 PDF 文件")

        out_dir = work_dir / "output"
        out_dir.mkdir()
        total_pages = 0
        unknown_pages: List[str] = []
        used_bases: Dict[str, int] = {}

        if group_by_sku:
            groups: "OrderedDict[str, fitz.Document]" = OrderedDict()
            qty_map: Dict[str, int] = OrderedDict()
            page_map: Dict[str, int] = OrderedDict()
            seller_map: Dict[str, List[str]] = OrderedDict()

            for pdf_path in pdf_paths:
                src = fitz.open(str(pdf_path))
                try:
                    for i, page in enumerate(src):
                        seller_sku, qty = extract_sku_and_qty(page)
                        if seller_sku == "UNKNOWN-SKU":
                            unknown_pages.append(f"{pdf_path.name} page {i + 1}")
                        display_sku = mapped_sku_for_naming(seller_sku, naming_mode, sku_mapping)
                        group_key = display_sku if naming_mode == "warehouse" else seller_sku
                        total_pages += 1
                        if group_key not in groups:
                            groups[group_key] = fitz.open()
                            qty_map[group_key] = 0
                            page_map[group_key] = 0
                            seller_map[group_key] = []
                        if seller_sku not in seller_map[group_key]:
                            seller_map[group_key].append(seller_sku)
                        one_page_doc = make_output_page(src, i, size_key, add_made, made_font_size)
                        groups[group_key].insert_pdf(one_page_doc)
                        one_page_doc.close()
                        qty_map[group_key] += qty
                        page_map[group_key] += 1
                finally:
                    src.close()

            manifest_lines = [
                "PDF label processing finished",
                "Output mode: group by SellerSKU",
                f"Input PDFs: {len(pdf_paths)}",
                f"Total pages: {total_pages}",
                f"SellerSKU groups: {len(groups)}",
                f"Naming mode: {'warehouse SKU' if naming_mode == 'warehouse' else 'SellerSKU'}",
                f"Mapping rows: {len(sku_mapping)}",
                f"Size: {size_key}",
                f"Made In China: {'yes' if add_made else 'no'}",
                "",
                "Groups:",
            ]
            skus_payload = []
            for group_key, doc in groups.items():
                display_sku = group_key
                base = f"{safe_name(display_sku)}-{qty_map[group_key]}只"
                out_pdf = write_group_pdf(doc, out_dir, base, used_bases)
                doc.close()
                rel = out_pdf.relative_to(out_dir)
                sellers = ",".join(seller_map.get(group_key, []))
                manifest_lines.append(f"{rel}\tseller_sku={sellers}\tname_sku={display_sku}\tpages={page_map[group_key]}\tqty={qty_map[group_key]}")
                skus_payload.append({"sku": sellers, "name_sku": display_sku, "qty": qty_map[group_key], "pages": page_map[group_key]})
        else:
            manifest_lines = [
                "PDF label processing finished",
                "Output mode: original files",
                f"Input PDFs: {len(pdf_paths)}",
                f"Naming mode: {'warehouse SKU' if naming_mode == 'warehouse' else 'SellerSKU'}",
                f"Mapping rows: {len(sku_mapping)}",
                f"Size: {size_key}",
                f"Made In China: {'yes' if add_made else 'no'}",
                "",
                "Files:",
            ]
            skus_payload = []
            used_names: Dict[str, int] = {}
            for pdf_path in pdf_paths:
                src = fitz.open(str(pdf_path))
                out_doc = fitz.open()
                qty_sum = 0
                pages = 0
                file_skus: List[str] = []
                first_sku = ""
                try:
                    for i, page in enumerate(src):
                        sku, qty = extract_sku_and_qty(page)
                        if sku != "UNKNOWN-SKU":
                            file_skus.append(sku)
                            if not first_sku:
                                first_sku = sku
                        qty_sum += qty
                        pages += 1
                        total_pages += 1
                        one_page_doc = make_output_page(src, i, size_key, add_made, made_font_size)
                        out_doc.insert_pdf(one_page_doc)
                        one_page_doc.close()
                finally:
                    src.close()

                unique_skus = sorted(set(file_skus))
                if len(unique_skus) == 1:
                    name_sku = mapped_sku_for_naming(unique_skus[0], naming_mode, sku_mapping)
                    base = f"{safe_name(name_sku)}-{qty_sum}只"
                else:
                    base = safe_name(pdf_path.stem) or "processed"
                out_pdf = write_group_pdf(out_doc, out_dir, base, used_names)
                out_doc.close()
                manifest_lines.append(f"{out_pdf.relative_to(out_dir)}\tpages={pages}\tqty={qty_sum}\tfirst_seller_sku={first_sku or 'UNKNOWN-SKU'}")

        if unknown_pages:
            manifest_lines += ["", "Unknown SKU pages:", *unknown_pages]
        (out_dir / "处理说明.txt").write_text("\n".join(manifest_lines), encoding="utf-8")

        if output_zip.exists():
            output_zip.unlink()
        with zipfile.ZipFile(output_zip, "w", zipfile.ZIP_DEFLATED) as zf:
            for p in sorted(out_dir.rglob("*")):
                zf.write(p, p.relative_to(out_dir))

        return {
            "input_pdfs": len(pdf_paths),
            "total_pages": total_pages,
            "mode": "group_by_sku" if group_by_sku else "original_files",
            "groups": len(skus_payload),
            "skus": skus_payload,
            "unknown_pages": unknown_pages,
        }


app = FastAPI(title="PDF Label Tool v20")
app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/", response_class=HTMLResponse)
def index() -> str:
    return Path("static/index.html").read_text(encoding="utf-8")


@app.post("/api/inspect")
async def inspect(
    file: UploadFile = File(...),
    naming_mode: str = Form("seller"),
    mapping_file: Optional[UploadFile] = File(None),
):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        in_path = tmp / safe_name(file.filename or "input.pdf")
        in_path.write_bytes(await file.read())
        mapping_path = None
        if mapping_file and mapping_file.filename:
            mapping_path = tmp / safe_name(mapping_file.filename)
            mapping_path.write_bytes(await mapping_file.read())
        try:
            sku_mapping = load_sku_mapping(mapping_path) if naming_mode == "warehouse" else {}
            pdfs = collect_pdfs(in_path, tmp)
            group_counts: Dict[str, int] = OrderedDict()
            group_pages: Dict[str, int] = OrderedDict()
            seller_map: Dict[str, List[str]] = OrderedDict()
            total_pages = 0
            for pdf_path in pdfs:
                doc = fitz.open(str(pdf_path))
                try:
                    for page in doc:
                        seller_sku, qty = extract_sku_and_qty(page)
                        display_sku = mapped_sku_for_naming(seller_sku, naming_mode, sku_mapping)
                        group_key = display_sku if naming_mode == "warehouse" else seller_sku
                        group_counts[group_key] = group_counts.get(group_key, 0) + qty
                        group_pages[group_key] = group_pages.get(group_key, 0) + 1
                        seller_map.setdefault(group_key, [])
                        if seller_sku not in seller_map[group_key]:
                            seller_map[group_key].append(seller_sku)
                        total_pages += 1
                finally:
                    doc.close()
            groups = []
            for group_key, qty in group_counts.items():
                groups.append({
                    "sku": ",".join(seller_map.get(group_key, [])),
                    "name_sku": group_key,
                    "qty": qty,
                    "pages": group_pages.get(group_key, 0),
                    "file_base": f"{safe_name(group_key)}-{qty}只",
                })
            return {"total_pages": total_pages, "groups": groups, "naming_mode": naming_mode, "mapping_rows": len(sku_mapping)}
        except Exception as e:
            return JSONResponse({"error": str(e)}, status_code=400)


@app.post("/api/process")
async def process(
    file: UploadFile = File(...),
    size: str = Form("10x8"),
    add_made: bool = Form(False),
    made_font_size: float = Form(8.0),
    group_by_sku: bool = Form(True),
    naming_mode: str = Form("seller"),
    mapping_file: Optional[UploadFile] = File(None),
):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        in_path = tmp / safe_name(file.filename or "input.pdf")
        in_path.write_bytes(await file.read())
        mapping_path = None
        if mapping_file and mapping_file.filename:
            mapping_path = tmp / safe_name(mapping_file.filename)
            mapping_path.write_bytes(await mapping_file.read())
        out_zip = tmp / "PDF标签处理结果.zip"
        try:
            process_file(in_path, out_zip, size, add_made, made_font_size, group_by_sku, naming_mode, mapping_path)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))
        final_path = Path(tempfile.gettempdir()) / f"pdf-label-result-{os.getpid()}-{abs(hash(str(out_zip)))}.zip"
        shutil.copy2(out_zip, final_path)
        return FileResponse(final_path, filename="PDF标签处理结果.zip", media_type="application/zip")
